"""Excel export / import built on openpyxl."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.config import REPORTS_DIR
from app.services.report_service import DATE, MONEY, NUMBER, PERCENT, ReportData
from app.utils.i18n import tr

HEADER_FILL = PatternFill("solid", fgColor="1E293B")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=14)
GROUP_FONT = Font(bold=True)
THIN = Side(style="thin", color="C7D2E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MONEY_FORMAT = "# ##0"
NUMBER_FORMAT = "# ##0.00"
PERCENT_FORMAT = "0.0"
DATE_FORMAT = "DD.MM.YYYY"

#: Column headers accepted by the estimate importer (lower-cased).
IMPORT_HEADERS = {
    "section": "section",
    "bo'lim": "section",
    "bolim": "section",
    "code": "code",
    "kod": "code",
    "name": "name",
    "nomi": "name",
    "category": "category",
    "kategoriya": "category",
    "unit": "unit",
    "birlik": "unit",
    "quantity": "quantity",
    "miqdor": "quantity",
    "price": "price",
    "narx": "price",
    "note": "note",
    "izoh": "note",
}


def _output_path(name: str) -> Path:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return REPORTS_DIR / f"{name}_{stamp}.xlsx"


def _number_format(kind: str) -> str | None:
    return {
        MONEY: MONEY_FORMAT,
        NUMBER: NUMBER_FORMAT,
        PERCENT: PERCENT_FORMAT,
        DATE: DATE_FORMAT,
    }.get(kind)


def render_report(data: ReportData, path: str | Path | None = None) -> Path:
    """Write a :class:`ReportData` table into a styled worksheet."""
    target = Path(path) if path else _output_path(data.key)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = data.key[:28] or "report"

    sheet.cell(row=1, column=1, value=data.title).font = TITLE_FONT
    if data.subtitle:
        sheet.cell(row=2, column=1, value=data.subtitle)
    header_row = 4

    for index, column in enumerate(data.columns, start=1):
        cell = sheet.cell(row=header_row, column=index, value=column.title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        sheet.column_dimensions[get_column_letter(index)].width = max(12, column.width * 13)

    for offset, row in enumerate(data.rows, start=header_row + 1):
        for index, column in enumerate(data.columns, start=1):
            value = row.get(column.key)
            if isinstance(value, str) and column.kind in (MONEY, NUMBER, PERCENT):
                value = None
            cell = sheet.cell(row=offset, column=index, value=value)
            fmt = _number_format(column.kind)
            if fmt:
                cell.number_format = fmt
            if row.get("_group"):
                cell.font = GROUP_FONT
            cell.border = BORDER

    if data.totals:
        total_row = header_row + len(data.rows) + 1
        sheet.cell(row=total_row, column=1, value=tr("total")).font = GROUP_FONT
        for index, column in enumerate(data.columns, start=1):
            if column.key in data.totals:
                cell = sheet.cell(row=total_row, column=index, value=data.totals[column.key])
                cell.font = GROUP_FONT
                cell.number_format = _number_format(column.kind) or MONEY_FORMAT

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    workbook.save(target)
    return target


def export_rows(
    rows: list[dict],
    columns: list[tuple[str, str]],
    title: str,
    path: str | Path | None = None,
) -> Path:
    """Export arbitrary rows given ``[(key, header)]`` column pairs."""
    target = Path(path) if path else _output_path("export")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "data"
    sheet.cell(row=1, column=1, value=title).font = TITLE_FONT
    for index, (_key, header) in enumerate(columns, start=1):
        cell = sheet.cell(row=3, column=index, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        sheet.column_dimensions[get_column_letter(index)].width = 18
    for offset, row in enumerate(rows, start=4):
        for index, (key, _header) in enumerate(columns, start=1):
            value = row.get(key)
            if isinstance(value, (date, datetime)):
                cell = sheet.cell(row=offset, column=index, value=value)
                cell.number_format = DATE_FORMAT
            else:
                sheet.cell(row=offset, column=index, value=value)
    workbook.save(target)
    return target


def estimate_template(path: str | Path | None = None) -> Path:
    """Write an empty import template for estimate lines."""
    target = Path(path) if path else _output_path("estimate_template")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "estimate"
    headers = ["Section", "Code", "Name", "Category", "Unit", "Quantity", "Price", "Note"]
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        sheet.column_dimensions[get_column_letter(index)].width = 20
    sheet.cell(row=2, column=1, value="Tayyorlov ishlari")
    sheet.cell(row=2, column=2, value="1.1")
    sheet.cell(row=2, column=3, value="Demontaj")
    sheet.cell(row=2, column=5, value="m2")
    sheet.cell(row=2, column=6, value=100)
    sheet.cell(row=2, column=7, value=35000)
    workbook.save(target)
    return target


def read_estimate_rows(path: str | Path) -> list[dict]:
    """Read estimate lines from an Excel file.

    The first row must contain headers; both English and Uzbek captions from
    :data:`IMPORT_HEADERS` are recognised.
    """
    workbook = load_workbook(Path(path), data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(value or "").strip().lower() for value in rows[0]]
    mapping = {
        index: IMPORT_HEADERS[name] for index, name in enumerate(header) if name in IMPORT_HEADERS
    }
    if not mapping:
        raise ValueError("Excel header not recognised")
    result: list[dict] = []
    for raw in rows[1:]:
        record: dict = {}
        for index, key in mapping.items():
            if index < len(raw):
                record[key] = raw[index]
        if not (record.get("name") or record.get("code")):
            continue
        record["quantity"] = _to_float(record.get("quantity"))
        record["price"] = _to_float(record.get("price"))
        result.append(record)
    return result


def _to_float(value) -> float:
    try:
        return float(str(value).replace(" ", "").replace(",", ".")) if value is not None else 0.0
    except (TypeError, ValueError):
        return 0.0

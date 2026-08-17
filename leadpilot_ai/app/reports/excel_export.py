"""Excel export/import helpers built on openpyxl."""

from __future__ import annotations

import logging
from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.crm import Lead
from app.models.enums import Channel, LeadStatus
from app.utils.dates import fmt_date, fmt_datetime
from app.utils.formatting import normalize_phone

logger = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="1E293B")
HEADER_FONT = Font(color="E5E9F0", bold=True, size=11)


def _write_sheet(
    workbook: Workbook,
    title: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
    *,
    first: bool = False,
) -> None:
    """Write one styled worksheet."""
    sheet = workbook.active if first else workbook.create_sheet()
    sheet.title = title[:31]
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        sheet.append(list(row))
    for index, header in enumerate(headers, start=1):
        width = max(12, min(42, len(str(header)) + 6))
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def export_table(
    path: str, sheets: list[tuple[str, Sequence[str], Iterable[Sequence[Any]]]]
) -> str:
    """Write one or more sheets to an ``.xlsx`` file and return its path."""
    workbook = Workbook()
    for index, (title, headers, rows) in enumerate(sheets):
        _write_sheet(workbook, title, headers, rows, first=index == 0)
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    logger.info("Excel export written: %s", target)
    return str(target)


def export_leads(leads: Iterable[Lead], path: str) -> str:
    """Export the lead registry."""
    headers = [
        "ID",
        "Ism",
        "Telefon",
        "Email",
        "Telegram",
        "Til",
        "Kanal",
        "Status",
        "Niyat",
        "Skor",
        "Operator",
        "Xizmat",
        "Filial",
        "Manba",
        "Kampaniya",
        "utm_source",
        "utm_medium",
        "utm_campaign",
        "Daromad",
        "Yo'qotish sababi",
        "Yaratilgan",
        "So'nggi faollik",
    ]
    rows = []
    for lead in leads:
        rows.append(
            [
                lead.id,
                lead.display_name,
                lead.phone or "",
                lead.email,
                lead.telegram_username,
                lead.language,
                lead.channel,
                lead.status,
                lead.intent,
                lead.score,
                lead.owner.full_name if lead.owner else "",
                lead.service.name if lead.service else "",
                lead.branch.name if lead.branch else "",
                lead.source.name if lead.source else lead.utm_source,
                lead.campaign.name if lead.campaign else lead.utm_campaign,
                lead.utm_source,
                lead.utm_medium,
                lead.utm_campaign,
                float(lead.revenue or 0),
                lead.loss_reason,
                fmt_datetime(lead.created_at),
                fmt_datetime(lead.last_activity_at),
            ]
        )
    return export_table(path, [("Leadlar", headers, rows)])


#: Column aliases accepted by the importer (lower-cased).
IMPORT_ALIASES: dict[str, list[str]] = {
    "full_name": ["ism", "name", "fio", "имя", "full_name"],
    "phone": ["telefon", "phone", "телефон", "raqam"],
    "email": ["email", "e-mail", "pochta"],
    "telegram_username": ["telegram", "username"],
    "channel": ["kanal", "channel", "канал"],
    "interest": ["qiziqish", "interest", "интерес", "xizmat"],
    "notes": ["izoh", "note", "комментарий", "notes"],
    "utm_source": ["utm_source"],
    "utm_medium": ["utm_medium"],
    "utm_campaign": ["utm_campaign"],
}


def import_leads(path: str) -> list[dict[str, str]]:
    """Read leads from an Excel file and return normalised payload dicts.

    The caller passes each dict to :func:`app.services.lead_service.create_lead`.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return []

    mapping: dict[int, str] = {}
    for index, raw in enumerate(header_row):
        if raw is None:
            continue
        header = str(raw).strip().lower()
        for field, aliases in IMPORT_ALIASES.items():
            if header in aliases:
                mapping[index] = field
                break

    result: list[dict[str, str]] = []
    for row in rows:
        payload: dict[str, str] = {}
        for index, value in enumerate(row):
            field = mapping.get(index)
            if not field or value in (None, ""):
                continue
            payload[field] = str(value).strip()
        if not payload:
            continue
        if "phone" in payload:
            payload["phone"] = normalize_phone(payload["phone"]) or payload["phone"]
        channel = payload.get("channel", "").lower()
        payload["channel"] = channel if channel in {c.value for c in Channel} else Channel.MANUAL
        payload["status"] = LeadStatus.NEW
        result.append(payload)
    workbook.close()
    logger.info("Excel import parsed %s rows from %s", len(result), path)
    return result


def export_generic(
    path: str, title: str, headers: Sequence[str], rows: Iterable[Sequence[Any]]
) -> str:
    """Convenience wrapper used by the reports page."""
    return export_table(path, [(title, headers, rows)])


def format_date_cell(value) -> str:  # type: ignore[no-untyped-def]
    """Format a date for a spreadsheet cell."""
    return fmt_date(value)

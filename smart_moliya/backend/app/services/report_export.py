import io

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.schemas.report import ReportSummary


def build_pdf(summary: ReportSummary) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph("Smart Moliya — Moliyaviy hisobot", styles["Title"]),
        Paragraph(f"{summary.start_date} — {summary.end_date}", styles["Normal"]),
        Spacer(1, 16),
        Paragraph(
            f"Jami daromad: {summary.total_income:,.0f} so'm &nbsp;|&nbsp; "
            f"Jami xarajat: {summary.total_expense:,.0f} so'm &nbsp;|&nbsp; "
            f"Sof natija: {summary.net:,.0f} so'm",
            styles["Heading3"],
        ),
        Spacer(1, 16),
        Paragraph("Kategoriya bo'yicha xarajatlar", styles["Heading2"]),
    ]

    category_data = [["Kategoriya", "Summa", "%"]] + [
        [item.category_name, f"{item.amount:,.0f}", f"{item.percent}%"] for item in summary.by_category
    ]
    category_table = Table(category_data, hAlign="LEFT")
    category_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B5E20")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
            ]
        )
    )
    elements.append(category_table)
    elements.append(Spacer(1, 16))
    elements.append(Paragraph("Kunlik Cash Flow", styles["Heading2"]))

    cash_flow_data = [["Sana", "Daromad", "Xarajat", "Sof"]] + [
        [str(item.date), f"{item.income:,.0f}", f"{item.expense:,.0f}", f"{item.net:,.0f}"]
        for item in summary.daily_cash_flow
    ]
    cash_flow_table = Table(cash_flow_data, hAlign="LEFT")
    cash_flow_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B5E20")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
            ]
        )
    )
    elements.append(cash_flow_table)

    doc.build(elements)
    return buffer.getvalue()


def build_excel(summary: ReportSummary) -> bytes:
    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Xulosa"
    summary_sheet.append(["Davr", f"{summary.start_date} — {summary.end_date}"])
    summary_sheet.append(["Jami daromad", float(summary.total_income)])
    summary_sheet.append(["Jami xarajat", float(summary.total_expense)])
    summary_sheet.append(["Sof natija", float(summary.net)])
    for row in summary_sheet.iter_rows(min_row=1, max_row=4, min_col=1, max_col=1):
        row[0].font = Font(bold=True)

    category_sheet = workbook.create_sheet("Kategoriyalar")
    category_sheet.append(["Kategoriya", "Summa", "Foiz"])
    for cell in category_sheet[1]:
        cell.font = Font(bold=True)
    for item in summary.by_category:
        category_sheet.append([item.category_name, float(item.amount), item.percent])

    cash_flow_sheet = workbook.create_sheet("Cash Flow")
    cash_flow_sheet.append(["Sana", "Daromad", "Xarajat", "Sof"])
    for cell in cash_flow_sheet[1]:
        cell.font = Font(bold=True)
    for item in summary.daily_cash_flow:
        cash_flow_sheet.append([str(item.date), float(item.income), float(item.expense), float(item.net)])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()

from datetime import date
from decimal import Decimal

from app.schemas.report import CategoryBreakdownItem, DailyCashFlowItem, ReportSummary
from app.services.report_export import build_excel, build_pdf


def _sample_summary() -> ReportSummary:
    return ReportSummary(
        start_date=date(2026, 7, 1),
        end_date=date(2026, 7, 12),
        total_income=Decimal("2000000"),
        total_expense=Decimal("850000"),
        net=Decimal("1150000"),
        by_category=[
            CategoryBreakdownItem(category_id="c1", category_name="Oziq ovqat", amount=Decimal("500000"), percent=58.8),
            CategoryBreakdownItem(category_id="c2", category_name="Transport", amount=Decimal("350000"), percent=41.2),
        ],
        daily_cash_flow=[
            DailyCashFlowItem(date=date(2026, 7, 1), income=Decimal("2000000"), expense=Decimal("100000"), net=Decimal("1900000")),
            DailyCashFlowItem(date=date(2026, 7, 2), income=Decimal("0"), expense=Decimal("50000"), net=Decimal("-50000")),
        ],
    )


def test_build_pdf_returns_valid_pdf_bytes():
    pdf_bytes = build_pdf(_sample_summary())
    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 500


def test_build_excel_returns_valid_xlsx_bytes():
    excel_bytes = build_excel(_sample_summary())
    # XLSX fayllari ZIP konteyner - "PK" magic bytes bilan boshlanadi
    assert excel_bytes.startswith(b"PK")
    assert len(excel_bytes) > 500


def test_build_excel_sheets_contain_expected_data():
    import io
    from openpyxl import load_workbook

    excel_bytes = build_excel(_sample_summary())
    workbook = load_workbook(io.BytesIO(excel_bytes))

    assert workbook.sheetnames == ["Xulosa", "Kategoriyalar", "Cash Flow"]
    category_sheet = workbook["Kategoriyalar"]
    assert category_sheet["A2"].value == "Oziq ovqat"
    assert category_sheet["B2"].value == 500000.0

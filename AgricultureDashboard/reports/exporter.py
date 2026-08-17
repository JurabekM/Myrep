"""Report generation and data export in every supported format."""
from __future__ import annotations

import json
import logging
from datetime import datetime
from pathlib import Path

import pandas as pd

from analytics import kpi as kpi_module
from config import settings
from core.optional import try_import
from core.security import audit
from database.engine import read_df

log = logging.getLogger(__name__)


def _stamp(name: str, extension: str) -> Path:
    settings.ensure_directories()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return settings.EXPORT_DIR / f"{name}_{timestamp}.{extension}"


def _latin(text: str) -> str:
    """Sanitise Uzbek-latin text for PDF core fonts (latin-1 only)."""
    replacements = {"ʻ": "'", "ʼ": "'", "‘": "'", "’": "'", "“": '"', "”": '"',
                    "–": "-", "—": "-", "²": "2", "³": "3"}
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text.encode("latin-1", errors="replace").decode("latin-1")


# --------------------------------------------------------------------------
# DataFrame exports
# --------------------------------------------------------------------------
def export_dataframe(df: pd.DataFrame, name: str, fmt: str,
                     username: str = "") -> Path:
    """Export any table as csv / xlsx / json."""
    fmt = fmt.lower()
    if fmt == "csv":
        path = _stamp(name, "csv")
        df.to_csv(path, index=False, encoding="utf-8-sig")
    elif fmt == "xlsx":
        path = _stamp(name, "xlsx")
        df.to_excel(path, index=False, engine="openpyxl")
    elif fmt == "json":
        path = _stamp(name, "json")
        path.write_text(df.to_json(orient="records", force_ascii=False, indent=2),
                        encoding="utf-8")
    else:
        raise ValueError(f"Noma'lum format: {fmt}")
    audit("export", username, f"{name}.{fmt} ({len(df)} qator)")
    return path


def export_fields_geojson(username: str = "") -> Path:
    """All field polygons + attributes as a GeoJSON FeatureCollection."""
    df = read_df(
        "SELECT f.name, f.area_ha, f.soil_type, f.geometry_geojson,"
        " fa.name AS farm, r.name AS region"
        " FROM fields f"
        " JOIN farms fa ON fa.id = f.farm_id"
        " JOIN districts d ON d.id = fa.district_id"
        " JOIN regions r ON r.id = d.region_id"
    )
    features = []
    for row in df.itertuples():
        try:
            geometry = json.loads(row.geometry_geojson)
        except (json.JSONDecodeError, TypeError):
            continue
        features.append({
            "type": "Feature", "geometry": geometry,
            "properties": {"name": row.name, "area_ha": row.area_ha,
                           "soil": row.soil_type, "farm": row.farm,
                           "region": row.region},
        })
    path = _stamp("dalalar", "geojson")
    path.write_text(json.dumps({"type": "FeatureCollection", "features": features},
                               ensure_ascii=False), encoding="utf-8")
    audit("export", username, f"GeoJSON ({len(features)} poligon)")
    return path


# --------------------------------------------------------------------------
# KPI summary reports
# --------------------------------------------------------------------------
def _report_tables(year: int) -> dict[str, pd.DataFrame]:
    return {
        "Viloyatlar bo'yicha hosildorlik": kpi_module.yield_by_region(year),
        "Ekinlar taqsimoti": kpi_module.crop_distribution(year).head(12),
        "Eng yaxshi xo'jaliklar": kpi_module.top_farms(year),
    }


def kpi_report_pdf(year: int | None = None, username: str = "") -> Path:
    """Multi-section PDF summary report (fpdf2, no external binaries)."""
    from fpdf import FPDF

    year = year or kpi_module.latest_year()
    kpi = kpi_module.compute_kpi(year)
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("helvetica", "B", 16)
    pdf.cell(0, 10, _latin(f"AgroVision — {year}-yil yakuniy hisoboti"),
             new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("helvetica", "", 9)
    pdf.cell(0, 6, _latin(f"Yaratildi: {datetime.now():%Y-%m-%d %H:%M}"),
             new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    pdf.set_font("helvetica", "B", 12)
    pdf.cell(0, 8, _latin("Asosiy ko'rsatkichlar"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("helvetica", "", 10)
    for label, value in [
        ("Umumiy yer maydoni", f"{kpi.total_area_ha:,.0f} ga"),
        ("Fermerlar / xo'jaliklar / dalalar",
         f"{kpi.farmers} / {kpi.farms} / {kpi.fields}"),
        ("O'rtacha hosildorlik", f"{kpi.avg_yield_t_ha} t/ga"),
        ("Jami ishlab chiqarish", f"{kpi.production_t:,.0f} t"),
        ("Daromad", f"{kpi.income / 1e9:,.2f} mlrd so'm"),
        ("Xarajat", f"{kpi.expense / 1e9:,.2f} mlrd so'm"),
        ("Sof foyda (ROI)", f"{kpi.profit / 1e9:,.2f} mlrd so'm ({kpi.roi_percent}%)"),
        ("O'rtacha NDVI", f"{kpi.avg_ndvi}"),
        ("Sug'orish suvi", f"{kpi.water_million_m3} mln m3"),
    ]:
        pdf.cell(90, 7, _latin(label))
        pdf.cell(0, 7, _latin(str(value)), new_x="LMARGIN", new_y="NEXT")

    for title, table in _report_tables(year).items():
        pdf.ln(4)
        pdf.set_font("helvetica", "B", 12)
        pdf.cell(0, 8, _latin(title), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("helvetica", "B", 8)
        columns = list(table.columns)
        width = max(24, int(190 / max(1, len(columns))))
        for column in columns:
            pdf.cell(width, 6, _latin(str(column))[:22], border=1)
        pdf.ln()
        pdf.set_font("helvetica", "", 8)
        for _, row in table.head(15).iterrows():
            for column in columns:
                value = row[column]
                text = f"{value:,.0f}" if isinstance(value, float) and abs(value) > 100 \
                    else str(value)
                pdf.cell(width, 6, _latin(text)[:22], border=1)
            pdf.ln()

    path = _stamp(f"hisobot_{year}", "pdf")
    pdf.output(str(path))
    audit("export", username, f"PDF hisobot {year}")
    return path


def kpi_report_html(year: int | None = None, username: str = "") -> Path:
    """Standalone HTML report (also serves as the printable dashboard snapshot)."""
    year = year or kpi_module.latest_year()
    kpi = kpi_module.compute_kpi(year)
    sections = "".join(
        f"<h2>{title}</h2>{table.to_html(index=False, border=0)}"
        for title, table in _report_tables(year).items()
    )
    html = f"""<!DOCTYPE html><html lang="uz"><head><meta charset="utf-8">
<title>AgroVision hisobot {year}</title>
<style>
 body{{font-family:Segoe UI,Arial,sans-serif;margin:32px;color:#1b2a1b}}
 h1{{color:#2e7d32}} h2{{color:#33691e;border-bottom:2px solid #aed581;padding-bottom:4px}}
 table{{border-collapse:collapse;width:100%;margin:12px 0;font-size:14px}}
 td,th{{padding:6px 10px;border-bottom:1px solid #dcedc8;text-align:left}}
 th{{background:#f1f8e9}}
 .kpi{{display:inline-block;background:#f1f8e9;border-radius:12px;padding:14px 22px;
       margin:6px;min-width:160px}}
 .kpi b{{display:block;font-size:22px;color:#2e7d32}}
</style></head><body>
<h1>AgroVision — {year}-yil yakuniy hisoboti</h1>
<p>Yaratildi: {datetime.now():%Y-%m-%d %H:%M}</p>
<div>
 <div class="kpi"><b>{kpi.total_area_ha:,.0f} ga</b>Yer maydoni</div>
 <div class="kpi"><b>{kpi.avg_yield_t_ha} t/ga</b>O'rtacha hosildorlik</div>
 <div class="kpi"><b>{kpi.production_t:,.0f} t</b>Ishlab chiqarish</div>
 <div class="kpi"><b>{kpi.profit / 1e9:,.1f} mlrd</b>Sof foyda (so'm)</div>
 <div class="kpi"><b>{kpi.roi_percent}%</b>ROI</div>
 <div class="kpi"><b>{kpi.avg_ndvi}</b>O'rtacha NDVI</div>
</div>
{sections}
</body></html>"""
    path = _stamp(f"hisobot_{year}", "html")
    path.write_text(html, encoding="utf-8")
    audit("export", username, f"HTML hisobot {year}")
    return path


def kpi_report_word(year: int | None = None, username: str = "") -> Path | None:
    """Word (.docx) report when python-docx is installed; else None."""
    if try_import("docx") is None:
        return None
    from docx import Document

    year = year or kpi_module.latest_year()
    kpi = kpi_module.compute_kpi(year)
    document = Document()
    document.add_heading(f"AgroVision — {year}-yil yakuniy hisoboti", level=0)
    document.add_paragraph(f"Yaratildi: {datetime.now():%Y-%m-%d %H:%M}")
    document.add_heading("Asosiy ko'rsatkichlar", level=1)
    for line in [
        f"Yer maydoni: {kpi.total_area_ha:,.0f} ga",
        f"O'rtacha hosildorlik: {kpi.avg_yield_t_ha} t/ga",
        f"Ishlab chiqarish: {kpi.production_t:,.0f} t",
        f"Sof foyda: {kpi.profit / 1e9:,.2f} mlrd so'm (ROI {kpi.roi_percent}%)",
    ]:
        document.add_paragraph(line, style="List Bullet")
    for title, table in _report_tables(year).items():
        document.add_heading(title, level=1)
        word_table = document.add_table(rows=1, cols=len(table.columns))
        word_table.style = "Light Grid Accent 3"
        for i, column in enumerate(table.columns):
            word_table.rows[0].cells[i].text = str(column)
        for _, row in table.head(15).iterrows():
            cells = word_table.add_row().cells
            for i, column in enumerate(table.columns):
                cells[i].text = str(row[column])
    path = _stamp(f"hisobot_{year}", "docx")
    document.save(str(path))
    audit("export", username, f"Word hisobot {year}")
    return path


def dashboard_png(year: int | None = None, username: str = "") -> Path | None:
    """Dashboard snapshot as PNG (matplotlib when available; else None)."""
    if try_import("matplotlib") is None:
        return None
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    year = year or kpi_module.latest_year()
    trend = kpi_module.yield_trend()
    regions = kpi_module.yield_by_region(year)
    crops = kpi_module.crop_distribution(year).head(8)

    figure, axes = plt.subplots(2, 2, figsize=(14, 9))
    figure.suptitle(f"AgroVision — {year}-yil ko'rsatkichlari", fontsize=15)

    axes[0][0].plot(trend["year"], trend["avg_yield"], marker="o", color="#2e7d32")
    axes[0][0].set_title("Hosildorlik trendi (t/ga)")
    axes[0][0].grid(alpha=0.3)

    axes[0][1].barh(regions["region"], regions["production"], color="#66bb6a")
    axes[0][1].set_title("Viloyatlar — ishlab chiqarish (t)")
    axes[0][1].invert_yaxis()

    axes[1][0].pie(crops["area"], labels=crops["crop"], autopct="%1.0f%%",
                   textprops={"fontsize": 8})
    axes[1][0].set_title("Ekin maydonlari taqsimoti")

    axes[1][1].bar(trend["year"].astype(str), trend["production"], color="#9ccc65")
    axes[1][1].set_title("Yillik ishlab chiqarish (t)")

    figure.tight_layout(rect=(0, 0, 1, 0.96))
    path = _stamp(f"dashboard_{year}", "png")
    figure.savefig(path, dpi=130)
    plt.close(figure)
    audit("export", username, f"PNG snapshot {year}")
    return path


DATASETS: dict[str, str] = {
    "Hosildorlik yozuvlari": (
        "SELECT y.year, r.name AS region, d.name AS district, fa.name AS farm,"
        " f.name AS field, c.name AS crop, y.area_ha, y.yield_t_ha, y.production_t"
        " FROM yield_records y"
        " JOIN crops c ON c.id = y.crop_id"
        " JOIN fields f ON f.id = y.field_id"
        " JOIN farms fa ON fa.id = f.farm_id"
        " JOIN districts d ON d.id = fa.district_id"
        " JOIN regions r ON r.id = d.region_id ORDER BY y.year DESC"
    ),
    "Bozor narxlari": (
        "SELECT m.date, c.name AS crop, m.price_per_kg, m.market_name, m.source"
        " FROM market_prices m JOIN crops c ON c.id = m.crop_id ORDER BY m.date DESC"
    ),
    "Ob-havo yozuvlari": (
        "SELECT w.date, d.name AS district, w.t_min, w.t_max,"
        " w.precipitation_mm, w.humidity, w.source"
        " FROM weather_records w JOIN districts d ON d.id = w.district_id"
        " ORDER BY w.date DESC LIMIT 20000"
    ),
    "Moliya yozuvlari": (
        "SELECT fr.year, fa.name AS farm, fr.category, fr.amount, fr.note"
        " FROM finance_records fr JOIN farms fa ON fa.id = fr.farm_id"
        " ORDER BY fr.year DESC"
    ),
    "Sug'orish yozuvlari": (
        "SELECT i.date, f.name AS field, i.water_m3, i.method"
        " FROM irrigation_records i JOIN fields f ON f.id = i.field_id"
        " ORDER BY i.date DESC LIMIT 20000"
    ),
    "NDVI indekslari": (
        "SELECT s.date, f.name AS field, s.ndvi, s.evi, s.source"
        " FROM satellite_indices s JOIN fields f ON f.id = s.field_id"
        " ORDER BY s.date DESC LIMIT 20000"
    ),
}


def dataset_frame(name: str) -> pd.DataFrame:
    """Load one of the named export datasets."""
    return read_df(DATASETS[name])

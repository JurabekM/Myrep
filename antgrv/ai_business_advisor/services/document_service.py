import fitz  # PyMuPDF
import docx
import openpyxl
import pytesseract
from PIL import Image
from pathlib import Path
from core.logger import app_logger

class DocumentService:
    """
    Handles reading and parsing various document formats (PDF, DOCX, XLSX, images).
    """

    @staticmethod
    def extract_text(file_path: str | Path) -> str:
        path = Path(file_path)
        if not path.exists():
            app_logger.error(f"File not found: {path}")
            return ""

        ext = path.suffix.lower()
        try:
            if ext == '.pdf':
                return DocumentService._extract_pdf(path)
            elif ext == '.docx':
                return DocumentService._extract_docx(path)
            elif ext == '.xlsx':
                return DocumentService._extract_xlsx(path)
            elif ext in ['.png', '.jpg', '.jpeg']:
                return DocumentService._extract_image(path)
            elif ext == '.txt':
                with open(path, 'r', encoding='utf-8') as f:
                    return f.read()
            else:
                app_logger.warning(f"Unsupported file format: {ext}")
                return ""
        except Exception as e:
            app_logger.error(f"Error extracting text from {path}: {e}")
            return ""

    @staticmethod
    def _extract_pdf(path: Path) -> str:
        text = []
        with fitz.open(path) as doc:
            for page in doc:
                text.append(page.get_text())
        return "\n".join(text)

    @staticmethod
    def _extract_docx(path: Path) -> str:
        doc = docx.Document(path)
        return "\n".join([para.text for para in doc.paragraphs])

    @staticmethod
    def _extract_xlsx(path: Path) -> str:
        wb = openpyxl.load_workbook(path, data_only=True)
        text = []
        for sheet in wb.worksheets:
            text.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                # Filter out None values and join row items
                row_text = " | ".join([str(cell) for cell in row if cell is not None])
                if row_text:
                    text.append(row_text)
        return "\n".join(text)

    @staticmethod
    def _extract_image(path: Path) -> str:
        # Requires Tesseract OCR installed on the system
        try:
            img = Image.open(path)
            return pytesseract.image_to_string(img)
        except Exception as e:
            app_logger.error(f"OCR failed for {path}. Ensure Tesseract is installed. Error: {e}")
            return ""

document_service = DocumentService()

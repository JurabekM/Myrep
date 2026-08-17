# -*- coding: utf-8 -*-
"""
Eksport servisi — jadval ma'lumotlarini PDF / Excel / CSV / Word / JSON
formatlarida ``exports/`` papkasiga chiqaradi.

Universal interfeys::

    path = exporter.export_rows("xlsx", "savdo_hisoboti",
                                "Savdo hisoboti", headers, rows)

Ixtiyoriy kutubxona o'rnatilmagan bo'lsa aniq o'zbekcha xato beriladi
(dastur yiqilmaydi).
"""
from __future__ import annotations

import csv
import json
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Sequence

from src.core.errors import UzERPError, ValidationError
from src.modules.base import BaseService

FORMATS = ("csv", "json", "xlsx", "pdf", "docx")


def _stringify(value: Any) -> str:
    """Katakcha qiymatini matnga aylantiradi (None -> bo'sh)."""
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return f"{value:,.2f}".replace(",", " ")
    return str(value)


class ExportService(BaseService):
    """Hisobot jadvallarini turli formatlarda eksport qiluvchi servis."""

    def __init__(self, container) -> None:
        super().__init__(container)
        self.exports_dir = Path(container.get("base_dir")) / "exports"

    # ------------------------------------------------------------------ #
    #  Umumiy kirish nuqtasi
    # ------------------------------------------------------------------ #

    def export_rows(self, fmt: str, filename_base: str, title: str,
                    headers: Sequence[str], rows: Sequence[Sequence[Any]],
                    user: dict | None = None) -> Path:
        """
        Jadvalni berilgan formatda eksport qiladi va fayl yo'lini qaytaradi.

        :param fmt: ``csv`` | ``json`` | ``xlsx`` | ``pdf`` | ``docx``
        """
        fmt = (fmt or "").lower().strip()
        if fmt not in FORMATS:
            raise ValidationError(f"Noma'lum eksport formati: {fmt}")
        self.exports_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_base = "".join(c for c in filename_base
                            if c.isalnum() or c in ("_", "-")) or "export"
        path = self.exports_dir / f"{safe_base}_{stamp}.{fmt}"

        writer = getattr(self, f"_write_{fmt}")
        writer(path, title, list(headers), [list(r) for r in rows])

        self.audit.log("export", user=user, entity="file",
                       details=f"{path.name} ({len(rows)} qator)")
        return path

    # ------------------------------------------------------------------ #
    #  Format yozuvchilari
    # ------------------------------------------------------------------ #

    @staticmethod
    def _write_csv(path: Path, title: str, headers: list[str],
                   rows: list[list[Any]]) -> None:
        """CSV (UTF-8 BOM bilan — Excel to'g'ri ochadi)."""
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.writer(fh, delimiter=";")
            writer.writerow(headers)
            for row in rows:
                writer.writerow([_stringify(v) for v in row])

    @staticmethod
    def _write_json(path: Path, title: str, headers: list[str],
                    rows: list[list[Any]]) -> None:
        """JSON: sarlavha + yozuvlar ro'yxati (dict ko'rinishida)."""
        data = {
            "title": title,
            "exported_at": datetime.now().isoformat(timespec="seconds"),
            "rows": [dict(zip(headers, [_stringify(v) for v in row]))
                     for row in rows],
        }
        path.write_text(json.dumps(data, ensure_ascii=False, indent=2),
                        encoding="utf-8")

    @staticmethod
    def _write_xlsx(path: Path, title: str, headers: list[str],
                    rows: list[list[Any]]) -> None:
        """Excel (openpyxl): qalin sarlavha, muzlatilgan qator, kengliklar."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError as exc:
            raise UzERPError(
                "Excel eksport uchun openpyxl kerak: pip install openpyxl"
            ) from exc

        wb = Workbook()
        ws = wb.active
        ws.title = (title or "Hisobot")[:30]

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="2D3748")
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        widths = [len(str(h)) for h in headers]
        for row in rows:
            values = []
            for idx, value in enumerate(row):
                if isinstance(value, Decimal):
                    value = float(value)
                values.append(value)
                if idx < len(widths):
                    widths[idx] = max(widths[idx], len(_stringify(value)))
            ws.append(values)

        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = \
                min(width + 3, 50)
        ws.freeze_panes = "A2"
        wb.save(path)

    def _write_pdf(self, path: Path, title: str, headers: list[str],
                   rows: list[list[Any]]) -> None:
        """PDF (fpdf2): unicode shrift (Windows Arial) bilan jadval."""
        try:
            from fpdf import FPDF
        except ImportError as exc:
            raise UzERPError(
                "PDF eksport uchun fpdf2 kerak: pip install fpdf2") from exc

        pdf = FPDF(orientation="L" if len(headers) > 6 else "P")
        pdf.set_auto_page_break(auto=True, margin=12)
        pdf.add_page()

        font_family = self._register_unicode_font(pdf)
        pdf.set_font(font_family, "B", 14)
        pdf.cell(0, 10, title, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font(font_family, "", 8)
        pdf.cell(0, 6, f"UzERP | {datetime.now():%Y-%m-%d %H:%M}",
                 new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)

        page_width = pdf.w - pdf.l_margin - pdf.r_margin
        col_width = page_width / max(len(headers), 1)

        pdf.set_font(font_family, "B", 8)
        pdf.set_fill_color(45, 55, 72)
        pdf.set_text_color(255, 255, 255)
        for header in headers:
            pdf.cell(col_width, 7, _stringify(header)[:40], border=1,
                     fill=True, align="C")
        pdf.ln()

        pdf.set_font(font_family, "", 8)
        pdf.set_text_color(0, 0, 0)
        fill = False
        pdf.set_fill_color(240, 242, 245)
        for row in rows:
            for value in row:
                pdf.cell(col_width, 6, _stringify(value)[:40], border=1,
                         fill=fill)
            pdf.ln()
            fill = not fill
        pdf.output(str(path))

    @staticmethod
    def _register_unicode_font(pdf) -> str:
        """
        Unicode TTF shriftni ulashga urinadi (kirill/lotin diakritikalari
        uchun); topilmasa standart helvetica qoladi.
        """
        candidates = [
            Path(r"C:\Windows\Fonts\arial.ttf"),
            Path(r"C:\Windows\Fonts\calibri.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        ]
        bold_candidates = [
            Path(r"C:\Windows\Fonts\arialbd.ttf"),
            Path(r"C:\Windows\Fonts\calibrib.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ]
        for regular, bold in zip(candidates, bold_candidates):
            if regular.exists():
                try:
                    pdf.add_font("UzFont", "", str(regular))
                    pdf.add_font("UzFont", "B",
                                 str(bold if bold.exists() else regular))
                    return "UzFont"
                except Exception:  # noqa: BLE001
                    continue
        return "helvetica"

    @staticmethod
    def _write_docx(path: Path, title: str, headers: list[str],
                    rows: list[list[Any]]) -> None:
        """Word (python-docx): sarlavha + jadval."""
        try:
            import docx
        except ImportError as exc:
            raise UzERPError(
                "Word eksport uchun python-docx kerak: "
                "pip install python-docx") from exc

        document = docx.Document()
        document.add_heading(title, level=1)
        document.add_paragraph(
            f"UzERP | {datetime.now():%Y-%m-%d %H:%M}")

        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Light Grid Accent 1"
        for idx, header in enumerate(headers):
            cell = table.rows[0].cells[idx]
            cell.text = _stringify(header)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
        for row in rows:
            cells = table.add_row().cells
            for idx, value in enumerate(row):
                if idx < len(cells):
                    cells[idx].text = _stringify(value)
        document.save(str(path))

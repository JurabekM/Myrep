# -*- coding: utf-8 -*-
"""
Import servisi — CSV / Excel / JSON fayllardan ma'lumot yuklash.

Qo'llab-quvvatlanadi: mahsulotlar, mijozlar, ta'minotchilar.
Mavjud yozuvlar yangilanadi (mahsulot — SKU bo'yicha, kontragent — nom
bo'yicha), yangilari qo'shiladi. Har bir qator xatosi alohida qayd etiladi —
bitta buzuq qator butun importni to'xtatmaydi.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from src.core.errors import UzERPError, ValidationError
from src.modules.base import BaseService

#: Ustun nomlari lug'ati: tashqi nom (kichik harfda) -> ichki maydon
_PRODUCT_ALIASES = {
    "name": "name", "nomi": "name", "nom": "name", "наименование": "name",
    "sku": "sku", "kod": "sku", "artikul": "sku", "код": "sku",
    "barcode": "barcode", "shtrix": "barcode", "shtrix-kod": "barcode",
    "штрихкод": "barcode",
    "unit": "unit", "birlik": "unit", "o'lchov": "unit", "ед": "unit",
    "cost_price": "cost_price", "tannarx": "cost_price",
    "sale_price": "sale_price", "narx": "sale_price", "sotish narxi": "sale_price",
    "цена": "sale_price",
    "vat_rate": "vat_rate", "qqs": "vat_rate",
    "min_stock": "min_stock", "min zaxira": "min_stock",
}
_PARTNER_ALIASES = {
    "name": "name", "nomi": "name", "nom": "name", "наименование": "name",
    "tin": "tin", "inn": "tin", "stir": "tin", "инн": "tin",
    "phone": "phone", "telefon": "phone", "тел": "phone",
    "email": "email", "e-mail": "email",
    "address": "address", "manzil": "address", "адрес": "address",
    "code": "code", "kod": "code",
}


class ImportService(BaseService):
    """Tashqi fayllardan ma'lumot import qiluvchi servis."""

    # ------------------------------------------------------------------ #
    #  Ochiq API
    # ------------------------------------------------------------------ #

    def import_products(self, file_path: str | Path,
                        user: dict | None = None) -> dict:
        """Mahsulotlarni import qiladi (SKU bo'yicha upsert)."""
        rows = self._read_rows(Path(file_path))
        products = self.container.get("products")
        created = updated = 0
        errors: list[str] = []

        for line_no, raw in enumerate(rows, start=2):
            try:
                data = self._map_row(raw, _PRODUCT_ALIASES)
                if not data.get("name"):
                    raise ValidationError("nom (name) ustuni bo'sh")
                existing = None
                if data.get("sku"):
                    existing = self.db.query_one(
                        "SELECT id FROM products WHERE sku = ?",
                        (str(data["sku"]).strip(),))
                if existing:
                    products.update(existing["id"], data, user)
                    updated += 1
                else:
                    products.create(data, user)
                    created += 1
            except Exception as exc:  # noqa: BLE001 - qatorlab hisobot
                errors.append(f"{line_no}-qator: {exc}")

        result = {"created": created, "updated": updated, "errors": errors}
        self.audit.log("import.products", user=user,
                       details=f"{Path(file_path).name}: +{created}/~{updated}, "
                               f"xato {len(errors)}")
        return result

    def import_customers(self, file_path: str | Path,
                         user: dict | None = None) -> dict:
        """Mijozlarni import qiladi (nom bo'yicha upsert)."""
        return self._import_partners(file_path, "customers", user)

    def import_suppliers(self, file_path: str | Path,
                         user: dict | None = None) -> dict:
        """Ta'minotchilarni import qiladi (nom bo'yicha upsert)."""
        return self._import_partners(file_path, "suppliers", user)

    # ------------------------------------------------------------------ #
    #  Ichki yordamchilar
    # ------------------------------------------------------------------ #

    def _import_partners(self, file_path: str | Path, table: str,
                         user: dict | None) -> dict:
        rows = self._read_rows(Path(file_path))
        service = self.container.get(table)  # customers | suppliers
        created = updated = 0
        errors: list[str] = []

        for line_no, raw in enumerate(rows, start=2):
            try:
                data = self._map_row(raw, _PARTNER_ALIASES)
                name = str(data.get("name") or "").strip()
                if not name:
                    raise ValidationError("nom (name) ustuni bo'sh")
                existing = self.db.query_one(
                    f"SELECT id FROM {table} WHERE LOWER(name) = ?",
                    (name.lower(),))
                if existing:
                    service.update(existing["id"], data, user)
                    updated += 1
                else:
                    service.create(data, user)
                    created += 1
            except Exception as exc:  # noqa: BLE001
                errors.append(f"{line_no}-qator: {exc}")

        result = {"created": created, "updated": updated, "errors": errors}
        self.audit.log(f"import.{table}", user=user,
                       details=f"{Path(file_path).name}: +{created}/~{updated}, "
                               f"xato {len(errors)}")
        return result

    @staticmethod
    def _map_row(raw: dict, aliases: dict[str, str]) -> dict:
        """Tashqi ustun nomlarini ichki maydonlarga o'giradi."""
        data: dict[str, Any] = {}
        for key, value in raw.items():
            field = aliases.get(str(key).strip().lower())
            if field and value not in (None, ""):
                data[field] = value
        return data

    def _read_rows(self, path: Path) -> list[dict]:
        """Fayl kengaytmasiga qarab qatorlarni dict ro'yxati sifatida o'qiydi."""
        if not path.exists():
            raise UzERPError(f"Fayl topilmadi: {path}")
        suffix = path.suffix.lower()
        if suffix == ".csv":
            return self._read_csv(path)
        if suffix == ".json":
            return self._read_json(path)
        if suffix in (".xlsx", ".xlsm"):
            return self._read_xlsx(path)
        raise ValidationError(
            f"Qo'llab-quvvatlanmaydigan format: {suffix} "
            "(CSV, JSON yoki XLSX yuklang)")

    @staticmethod
    def _read_csv(path: Path) -> list[dict]:
        """CSV: ajratkich (`;` yoki `,`) avtomatik aniqlanadi."""
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        sample = text[:2048]
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        reader = csv.DictReader(text.splitlines(), delimiter=delimiter)
        return [dict(row) for row in reader]

    @staticmethod
    def _read_json(path: Path) -> list[dict]:
        """JSON: ro'yxat yoki {"rows": [...]} strukturasi."""
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, dict):
            data = data.get("rows") or data.get("items") or []
        if not isinstance(data, list):
            raise ValidationError("JSON strukturasi ro'yxat bo'lishi kerak.")
        return [row for row in data if isinstance(row, dict)]

    @staticmethod
    def _read_xlsx(path: Path) -> list[dict]:
        """Excel: birinchi varaq, birinchi qator — sarlavha."""
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise UzERPError(
                "Excel import uchun openpyxl kerak: "
                "pip install openpyxl") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            headers = [str(h or "").strip() for h in next(rows_iter)]
        except StopIteration:
            return []
        result = [dict(zip(headers, row)) for row in rows_iter]
        wb.close()
        return result

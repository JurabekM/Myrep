"""Generic table export to Excel and PDF used by every report."""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, Spacer, Table

from app.reports.pdf_base import footer_factory, make_document, styles, table_style
from app.utils.formatting import fmt_date, fmt_datetime, now

HEADER_FILL = PatternFill("solid", fgColor="1D4ED8")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _cell_value(value: Any) -> Any:
    """Convert a python value into something openpyxl can store."""
    if isinstance(value, dt.datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, (dt.date, str, int, float, bool)) or value is None:
        return value
    return str(value)


def export_excel(
    path: str | Path,
    title: str,
    columns: list[tuple[str, str]],
    rows: list[dict],
    *,
    meta: dict[str, str] | None = None,
    sheet_name: str = "Report",
) -> str:
    """Write rows to a formatted ``.xlsx`` file.

    ``columns`` is a list of ``(key, header)`` pairs.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name[:31]

    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, size=14, color="111827")
    sheet["A2"] = f"ExportFlow · {fmt_datetime(now())}"
    sheet["A2"].font = Font(size=9, color="6B7280")
    start_row = 4
    if meta:
        for offset, (key, value) in enumerate(meta.items()):
            sheet.cell(row=start_row + offset, column=1, value=key).font = Font(
                size=9, color="6B7280"
            )
            sheet.cell(row=start_row + offset, column=2, value=value).font = Font(size=9, bold=True)
        start_row += len(meta) + 1

    for index, (_key, header) in enumerate(columns, 1):
        cell = sheet.cell(row=start_row, column=index, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_index, row in enumerate(rows, start_row + 1):
        for col_index, (key, _header) in enumerate(columns, 1):
            cell = sheet.cell(row=row_index, column=col_index, value=_cell_value(row.get(key)))
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            if isinstance(row.get(key), dt.date) and not isinstance(row.get(key), dt.datetime):
                cell.number_format = "DD.MM.YYYY"

    for index, (key, header) in enumerate(columns, 1):
        longest = max(
            [len(str(header))] + [len(str(row.get(key) or "")) for row in rows[:400]] or [10]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 3, 11), 48)
    sheet.freeze_panes = sheet.cell(row=start_row + 1, column=1)
    sheet.auto_filter.ref = (
        f"A{start_row}:{get_column_letter(len(columns))}{start_row + max(len(rows), 1)}"
    )

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(path))
    return str(path)


def export_pdf(
    path: str | Path,
    title: str,
    columns: list[tuple[str, str]],
    rows: list[dict],
    *,
    meta: dict[str, str] | None = None,
    company_name: str = "",
    landscape_mode: bool = True,
) -> str:
    """Write rows to a landscape PDF table."""
    st = styles()
    doc = make_document(path, title, landscape_mode=landscape_mode)
    story: list = [Paragraph(title, st["title"])]
    subtitle = f"{company_name} · {fmt_datetime(now())}".strip(" ·")
    story.append(Paragraph(subtitle, st["subtitle"]))
    if meta:
        story.append(
            Paragraph(
                " &nbsp;|&nbsp; ".join(f"<b>{key}:</b> {value}" for key, value in meta.items()),
                st["small"],
            )
        )
        story.append(Spacer(1, 4 * mm))

    header = [header for _key, header in columns]
    data = [header]
    for row in rows:
        line = []
        for key, _header in columns:
            value = row.get(key)
            if isinstance(value, dt.datetime):
                line.append(fmt_datetime(value))
            elif isinstance(value, dt.date):
                line.append(fmt_date(value))
            elif isinstance(value, float):
                line.append(f"{value:,.2f}")
            elif value is None:
                line.append("")
            else:
                line.append(Paragraph(str(value), st["cell"]))
        data.append(line)

    if len(data) == 1:
        data.append(["—"] * len(header))

    table = Table(data, repeatRows=1)
    table.setStyle(table_style())
    story.append(table)
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(f"Rows: {len(rows)}", st["small"]))

    footer = footer_factory(company_name or "ExportFlow", f"| {title}")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return str(path)


def read_excel_rows(path: str | Path, expected: list[str] | None = None) -> list[dict]:
    """Read the first sheet of a workbook into dictionaries keyed by header."""
    workbook = load_workbook(str(path), data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    header: list[str] = []
    for raw in rows_iter:
        candidate = [str(value).strip() if value is not None else "" for value in raw]
        if any(candidate):
            header = candidate
            break
    if expected:
        missing = [name for name in expected if name not in header]
        if missing:
            raise ValueError(f"Missing columns: {', '.join(missing)}")
    result: list[dict] = []
    for raw in rows_iter:
        if raw is None or not any(value is not None and str(value).strip() for value in raw):
            continue
        result.append({header[i]: raw[i] for i in range(min(len(header), len(raw)))})
    return result

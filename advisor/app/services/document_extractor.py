"""Yuklangan fayllardan matn ajratish (lokal Python kutubxonalari).

Barcha og'ir kutubxonalar lazy import qilinadi — o'rnatilmagan bo'lsa,
tushunarli xato beriladi (ilova baribir ishga tushadi).
"""

from __future__ import annotations

import csv
import io

from app.core.exceptions import DocumentError

MAX_EXTRACT_CHARS = 300_000


def extract_text(filename: str, data: bytes) -> str:
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    extractor = _EXTRACTORS.get(ext)
    if extractor is None:
        raise DocumentError(f"Qo'llab-quvvatlanmaydigan fayl turi: {ext}")
    text = extractor(data)
    if not text.strip():
        raise DocumentError("Fayldan matn topilmadi (bo'sh yoki skanlangan bo'lishi mumkin)")
    return text[:MAX_EXTRACT_CHARS]


def _require(module: str, pip_name: str):  # type: ignore[no-untyped-def]
    try:
        return __import__(module)
    except ImportError as exc:
        raise DocumentError(
            f"'{pip_name}' o'rnatilmagan. O'rnatish: pip install {pip_name}"
        ) from exc


def _extract_pdf(data: bytes) -> str:
    _require("pypdf", "pypdf")
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if len(text.strip()) < 50:
        # Skanlangan PDF bo'lishi mumkin — OCR sinaymiz.
        ocr = _ocr_pdf(reader)
        if ocr.strip():
            return ocr
    return text


def _ocr_pdf(reader) -> str:  # type: ignore[no-untyped-def]
    texts: list[str] = []
    for page in reader.pages:
        try:
            for image in page.images:
                texts.append(_ocr_image(image.data))
        except Exception:
            continue
    return "\n\n".join(t for t in texts if t.strip())


def _extract_docx(data: bytes) -> str:
    _require("docx", "python-docx")
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    _require("openpyxl", "openpyxl")
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"## {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    workbook.close()
    return "\n".join(parts)


def _extract_txt(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_csv(data: bytes) -> str:
    text = _extract_txt(data)
    try:
        dialect = csv.Sniffer().sniff(text[:2000])
    except csv.Error:
        dialect = csv.excel
    return "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(text), dialect))


def _ocr_image(data: bytes) -> str:
    _require("PIL", "Pillow")
    _require("pytesseract", "pytesseract")
    import pytesseract
    from PIL import Image

    image = Image.open(io.BytesIO(data))
    try:
        return pytesseract.image_to_string(image, lang="uzb+rus+eng")
    except pytesseract.TesseractError:
        return pytesseract.image_to_string(image, lang="eng")
    except Exception as exc:
        raise DocumentError(
            "OCR uchun Tesseract dasturi o'rnatilmagan bo'lishi mumkin "
            "(https://github.com/UB-Mannheim/tesseract/wiki)"
        ) from exc


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
    ".txt": _extract_txt,
    ".csv": _extract_csv,
    ".png": _ocr_image,
    ".jpg": _ocr_image,
    ".jpeg": _ocr_image,
}

SUPPORTED_EXTENSIONS = tuple(_EXTRACTORS.keys())

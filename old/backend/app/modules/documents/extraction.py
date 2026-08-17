"""Text extraction from uploaded files (sync — runs inside Celery workers)."""

import csv
import io

from app.core.exceptions import UnsupportedFileTypeError, ValidationError

MAX_EXTRACT_CHARS = 400_000


def extract_text(filename: str, data: bytes) -> str:
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    extractor = _EXTRACTORS.get(ext)
    if extractor is None:
        raise UnsupportedFileTypeError(
            "Bu fayl turidan matn ajratib bo'lmaydi", details={"extension": ext}
        )
    text = extractor(data)
    if not text.strip():
        raise ValidationError("Fayldan matn topilmadi (bo'sh yoki skanlangan bo'lishi mumkin)")
    return text[:MAX_EXTRACT_CHARS]


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n\n".join(pages)
    if len(text.strip()) < 50 and len(reader.pages) > 0:
        # Likely a scanned PDF — try OCR page images.
        return _ocr_pdf(data)
    return text


def _ocr_pdf(data: bytes) -> str:
    """OCR fallback for scanned PDFs via embedded images."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    texts: list[str] = []
    for page in reader.pages:
        for image in page.images:
            texts.append(_ocr_image(image.data))
    return "\n\n".join(t for t in texts if t.strip())


def _extract_docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"## {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    workbook.close()
    return "\n".join(parts)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for i, slide in enumerate(presentation.slides, start=1):
        parts.append(f"## Slayd {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def _extract_txt(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_csv(data: bytes) -> str:
    text = _extract_txt(data)
    try:
        dialect = csv.Sniffer().sniff(text[:2000])
    except csv.Error:
        dialect = csv.excel
    rows = csv.reader(io.StringIO(text), dialect)
    return "\n".join(" | ".join(row) for row in rows)


def _ocr_image(data: bytes) -> str:
    import pytesseract
    from PIL import Image

    image = Image.open(io.BytesIO(data))
    # uz (lotin) + rus + eng — mavjud bo'lmagan til paketi bo'lsa eng-ga tushadi
    try:
        return pytesseract.image_to_string(image, lang="uzb+rus+eng")
    except pytesseract.TesseractError:
        return pytesseract.image_to_string(image, lang="eng")


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
    ".pptx": _extract_pptx,
    ".txt": _extract_txt,
    ".csv": _extract_csv,
    ".png": _ocr_image,
    ".jpg": _ocr_image,
    ".jpeg": _ocr_image,
}
